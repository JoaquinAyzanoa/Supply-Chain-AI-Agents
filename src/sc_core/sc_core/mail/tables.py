"""Tables inside mail attachments (CSV and Excel), read in memory.

Suppliers send price lists as spreadsheets. This module turns such an
attachment into rows of cell strings and nothing more; what a price list
means is decided by the agent that reads it. Like the PDF helper, a bad or
huge file never raises: it yields no table and logs why.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from loguru import logger
from pydantic import Field

from sc_core.mail.models import Attachment
from sc_core.schema.base import StrictModel

DEFAULT_MAX_BYTES = 5_000_000
DEFAULT_MAX_ROWS = 2_000
SPREADSHEET_TYPES = {
    "text/csv",
    "application/csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}


class TableData(StrictModel):
    """One sheet (or one CSV) as rows of strings; empty cells are empty strings."""

    name: str
    rows: list[list[str]] = Field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return not any(any(cell for cell in row) for row in self.rows)


def is_spreadsheet(attachment: Attachment) -> bool:
    lower = attachment.name.lower()
    return attachment.content_type in SPREADSHEET_TYPES or lower.endswith(
        (".csv", ".xlsx", ".xlsm")
    )


def read_tables(
    attachment: Attachment,
    *,
    max_bytes: int = DEFAULT_MAX_BYTES,
    max_rows: int = DEFAULT_MAX_ROWS,
) -> list[TableData]:
    """The tables in a CSV (one) or an Excel workbook (one per sheet with content)."""
    if not is_spreadsheet(attachment) or not attachment.data:
        return []
    if len(attachment.data) > max_bytes:
        logger.bind(name=attachment.name, size=len(attachment.data)).warning(
            "spreadsheet skipped: larger than the limit"
        )
        return []
    lower = attachment.name.lower()
    try:
        if lower.endswith(".csv") or attachment.content_type in ("text/csv", "application/csv"):
            return [_csv(attachment.name, attachment.data, max_rows)]
        return _xlsx(attachment.name, attachment.data, max_rows)
    except Exception as exc:  # noqa: BLE001 - one odd file must not stop the message
        logger.bind(name=attachment.name).warning("spreadsheet unreadable: {}", exc)
        return []


def _csv(name: str, data: bytes, max_rows: int) -> TableData:
    text = _decode(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: list[list[str]] = []
    for row in csv.reader(io.StringIO(text), dialect):
        rows.append([cell.strip() for cell in row])
        if len(rows) >= max_rows:
            break
    return TableData(name=name, rows=rows)


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _xlsx(name: str, data: bytes, max_rows: int) -> list[TableData]:
    from openpyxl import load_workbook  # imported here: openpyxl is slow to import

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables: list[TableData] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for values in sheet.iter_rows(values_only=True):
                rows.append([_cell(v) for v in values])
                if len(rows) >= max_rows:
                    break
            table = TableData(name=f"{name}:{sheet.title}", rows=rows)
            if not table.is_empty:
                tables.append(table)
    finally:
        workbook.close()
    return tables


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


__all__ = ["TableData", "is_spreadsheet", "read_tables"]
